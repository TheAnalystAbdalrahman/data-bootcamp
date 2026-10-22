"""Read-only metadata catalog of the official course workbooks."""

from collections import Counter
from pathlib import Path
import re
import warnings
from urllib.parse import quote

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
OUT = Path(__file__).with_name("WORKBOOK_CATALOG.md")

SKILLS = {
    1: ["Excel interface and file handling", "Navigation and selection", "View options", "Data entry and types", "Fill handle"],
    2: ["Basic formulas", "Business calculations", "SUM and AutoSum", "AVERAGE, MIN and MAX", "Absolute references", "Cross-sheet calculations"],
    3: ["Font formatting", "Borders", "Alignment", "Format Painter", "Number formats", "Styles and themes"],
    4: ["Managing rows and columns", "Find and Replace", "Filtering", "Sorting", "Conditional formatting"],
    5: ["Print preview", "Orientation, margins and scaling", "Page breaks", "Repeating print titles", "Headers and footers"],
    6: ["Basic chart types", "Moving and resizing charts", "Chart style and type", "Chart elements"],
}


def esc(value):
    return str(value).replace("|", "\\|").replace("\n", " ")


def classify(path, week):
    name = path.stem.lower()
    if "soln" in name or "solution" in name:
        return "Completed example (filename)"
    if week == 4:
        return "Practice/data workbook"
    return "Exercise workbook"


files = sorted(ROOT.glob("[0-9][0-9]_*/*/**/*.xlsx"))
assert len(files) == 45, f"Expected 45 XLSX files, found {len(files)}"
rows = []
counts = Counter()
for path in files:
    week = int(path.name[1:3])
    video = int(re.search(r"V(\d\d)", path.name, re.I).group(1))
    counts[week] += 1
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    sheet_names = [ws.title for ws in wb.worksheets]
    funcs = Counter()
    formula_count = 0
    cross_sheet = 0
    features = []
    chart_count = sum(len(ws._charts) for ws in wb.worksheets)
    cf_count = sum(len(ws.conditional_formatting) for ws in wb.worksheets)
    table_count = sum(len(ws.tables) for ws in wb.worksheets)
    validation_count = sum(len(ws.data_validations.dataValidation) for ws in wb.worksheets)
    filter_count = sum(bool(ws.auto_filter and ws.auto_filter.ref) for ws in wb.worksheets)
    merged_count = sum(len(ws.merged_cells.ranges) for ws in wb.worksheets)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" and isinstance(cell.value, str):
                    formula_count += 1
                    cross_sheet += "!" in cell.value
                    funcs.update(x.upper() for x in re.findall(r"\b([A-Z][A-Z0-9_.]*)\s*\(", cell.value, re.I))
    if formula_count:
        features.append(f"{formula_count} formula cell{'s' if formula_count != 1 else ''}")
    if funcs:
        features.append("functions: " + ", ".join(sorted(funcs)))
    elif formula_count:
        features.append("arithmetic/cell references")
    if cross_sheet:
        features.append(f"{cross_sheet} cross-sheet formula{'s' if cross_sheet != 1 else ''}")
    if chart_count:
        features.append(f"{chart_count} chart{'s' if chart_count != 1 else ''}")
    if cf_count:
        features.append(f"{cf_count} conditional-format rule{'s' if cf_count != 1 else ''}")
    if table_count:
        features.append(f"{table_count} Excel table{'s' if table_count != 1 else ''}")
    if validation_count:
        features.append(f"{validation_count} data-validation rule{'s' if validation_count != 1 else ''}")
    if filter_count:
        features.append(f"{filter_count} filter range{'s' if filter_count != 1 else ''}")
    if merged_count:
        features.append(f"{merged_count} merged range{'s' if merged_count != 1 else ''}")
    if not features:
        features.append("No formulas or listed structural features found")
    rel = path.relative_to(ROOT)
    link = quote("../" + rel.as_posix(), safe="/._-")
    rows.append((week, video, path.name, link, sheet_names, SKILLS[week][video - 1], features, classify(path, week)))
    wb.close()

lines = [
    "# Official workbook catalog",
    "",
    "Read-only inspection of the 45 extracted `.xlsx` files on 17 October 2026. Original workbooks were not edited. Skill and role labels are inferred from filenames and structure, not claims about lesson content. Formula/function and feature counts describe the files as stored; they do not evaluate correctness.",
    "",
    "## Find a practice workbook by skill",
    "",
]
for week in range(1, 7):
    relevant = [r for r in rows if r[0] == week and "Completed" not in r[7]]
    listing = "; ".join(f"{r[5]}: [{r[2]}]({r[3]})" for r in relevant)
    lines.append(f"- **Module {week}:** {listing}.")
lines += [
    "",
    "## File-by-file inventory",
    "",
    "| Module | Original filename | Worksheets | Apparent skill | Formula/function and workbook features | Apparent role |",
    "| ---: | --- | --- | --- | --- | --- |",
]
for week, video, name, link, sheets, skill, features, role in rows:
    lines.append(f"| {week} | [{esc(name)}]({link}) | {esc(', '.join(sheets))} | {esc(skill)} | {esc('; '.join(features))} | {esc(role)} |")
lines += [
    "",
    "## How to use this catalog",
    "",
    "For each skill, open the non-solution workbook first and work on a **copy** so the official download remains unchanged. Where a matching `Soln` file exists, compare it only after trying the exercise. Modules 4 and 6 have no separately named solution files in the downloaded ZIPs; treat their workbooks as practice data, not answer keys. No file is explicitly named a template or dataset-only file, so those roles were not asserted. The ZIP archives remain in each module's `02_Official_Downloads` folder.",
    "",
    "Workbook counts: " + ", ".join(f"Module {week}: {counts[week]}" for week in range(1, 7)) + ". Total: 45.",
    "",
]
OUT.write_text("\n".join(lines), encoding="utf-8")
print(f"Cataloged {len(rows)} workbooks; counts: {dict(sorted(counts.items()))}")
